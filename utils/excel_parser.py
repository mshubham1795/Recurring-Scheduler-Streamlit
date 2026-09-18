"""
Excel/CSV parsing and template generation using openpyxl.
Replaces the SheetJS (XLSX) library used in the original frontend.
"""
import time
import logging
from io import BytesIO

import openpyxl
from openpyxl.utils import get_column_letter

from utils.time_utils import fix_time, fix_date
from utils.validators import validate_time, validate_path_no_spaces, validate_end_date_not_expired

log = logging.getLogger("SASBackend")


def parse_upload(uploaded_file):
    """Parse uploaded Excel/CSV into list of job dicts.

    Args:
        uploaded_file: Streamlit UploadedFile object

    Returns:
        list of job dicts with validation flags
    """
    filename = uploaded_file.name.lower()

    if filename.endswith('.csv'):
        return _parse_csv(uploaded_file)
    else:
        return _parse_excel(uploaded_file)


def _parse_csv(uploaded_file):
    """Parse a CSV file into job list."""
    import csv
    import io

    content = uploaded_file.read().decode("utf-8", errors="replace")
    uploaded_file.seek(0)
    reader = csv.DictReader(io.StringIO(content))
    rows = list(reader)
    return _rows_to_jobs(rows)


def _parse_excel(uploaded_file):
    """Parse an Excel file into job list."""
    wb = openpyxl.load_workbook(uploaded_file, data_only=True, read_only=True)

    # Find "Schedule" sheet or use last sheet
    sheet_name = None
    for sn in wb.sheetnames:
        if sn.lower() == "schedule":
            sheet_name = sn
            break
    if not sheet_name:
        sheet_name = wb.sheetnames[-1]

    ws = wb[sheet_name]

    # Read headers from first row
    rows_data = []
    headers = None
    for row in ws.iter_rows(values_only=True):
        if headers is None:
            headers = [str(cell).strip() if cell else "" for cell in row]
            continue
        if all(cell is None for cell in row):
            continue
        row_dict = {}
        for i, cell in enumerate(row):
            if i < len(headers) and headers[i]:
                row_dict[headers[i]] = str(cell).strip() if cell is not None else ""
        if any(v for v in row_dict.values()):
            rows_data.append(row_dict)

    wb.close()
    return _rows_to_jobs(rows_data)


def _rows_to_jobs(rows):
    """Convert raw row dicts to validated job dicts."""
    if not rows:
        return []

    # Validate required columns (case-insensitive)
    required = ['Study_Name', 'File_Path', 'File_Name', 'Start_Time', 'Frequency']
    cols = list(rows[0].keys())

    def get_col(row, target):
        """Case-insensitive column lookup."""
        for key in row:
            if key.strip().lower() == target.lower():
                return str(row[key]).strip()
        return ""

    # Check for missing required columns
    missing = []
    for req in required:
        found = any(c.strip().lower() == req.lower() for c in cols)
        if not found:
            missing.append(req)
    if missing:
        raise ValueError(f"Missing required columns: {', '.join(missing)}")

    jobs = []
    for i, row in enumerate(rows):
        st = fix_time(get_col(row, 'Start_Time'))
        ed = get_col(row, 'End_Date')
        if ed:
            ed = fix_date(ed)

        path = get_col(row, 'File_Path')
        file_name = get_col(row, 'File_Name')

        job = {
            "id": f"j{i}_{int(time.time() * 1000)}",
            "study": get_col(row, 'Study_Name'),
            "path": path,
            "file": file_name,
            "time": st or "",
            "freq": get_col(row, 'Frequency') or "Daily",
            "days": get_col(row, 'Days_of_Week') or "",
            "endDate": ed or "",
            "priority": get_col(row, 'Priority') or "Medium",
            "selected": False,
            "timeWarn": not validate_time(st) if st else True,
            "spaceWarn": not validate_path_no_spaces(path, file_name),
            "expiredWarn": not validate_end_date_not_expired(ed) if ed else False,
        }
        jobs.append(job)

    return jobs


def create_template_workbook():
    """Create an Excel template workbook with Instructions and Schedule sheets.

    Returns:
        BytesIO buffer containing the .xlsx file
    """
    wb = openpyxl.Workbook()

    # Instructions sheet
    ws_instr = wb.active
    ws_instr.title = "Instructions"

    instructions = [
        ["Recurring Scheduler -- Excel Template Instructions"],
        [""],
        ["REQUIRED COLUMNS:"],
        ["  Study_Name    -- Unique study identifier (e.g., GZQD, CPMP)"],
        ["  File_Path     -- Full directory path to the SAS program"],
        ["  File_Name     -- SAS script filename (must end in .sas)"],
        ["  Start_Time    -- Execution time in HH:MM 24-hour format"],
        ["  Frequency     -- Daily (Mon-Fri) | Weekly (specific days)"],
        [""],
        ["OPTIONAL COLUMNS:"],
        ["  Days_of_Week  -- Comma-separated days (e.g., Mon, Wed, Fri)"],
        ["                   Required when Frequency = Weekly"],
        ["  End_Date      -- When scheduling stops (YYYY-MM-DD). Leave blank for no end date."],
        ["  Priority      -- High | Medium | Low"],
        [""],
        ["HOW IT WORKS:"],
        ["  1. Fill in the Schedule sheet with your SAS programs"],
        ["  2. Upload this file in the app and click Save All"],
        ["  3. The backend will automatically submit jobs to CLUWE at the defined time/day"],
        ["  4. Jobs run until End_Date (if set), or forever if left blank"],
        ["  5. You can Run Now, Edit, Deactivate, or Reactivate schedules from the app"],
        [""],
        ["FREQUENCY OPTIONS:"],
        ["  Daily   -- Runs Monday to Friday at the specified time"],
        ["  Weekly  -- Runs only on specified days (set in Days_of_Week column)"],
        ["  Run Now -- Use in the app to execute a task immediately (not saved)"],
        [""],
        ["NOTES:"],
        ["  - File paths can use Z:\\ or /lillyce/ format"],
        ["  - Paths and filenames must NOT contain spaces"],
        ["  - Same file at different times = separate schedules"],
        ["  - CLUWE sends email notifications when jobs complete"],
    ]

    for row in instructions:
        ws_instr.append(row)
    ws_instr.column_dimensions['A'].width = 80

    # Schedule sheet
    ws_sched = wb.create_sheet("Schedule")

    headers = ['Study_Name', 'File_Path', 'File_Name', 'Start_Time',
               'Frequency', 'Days_of_Week', 'End_Date', 'Priority']
    ws_sched.append(headers)

    # Example data
    ws_sched.append([
        'GZQD',
        'Z:\\qa\\ly3537031\\j2s_mc_gzqd\\diva\\programs\\oversight',
        'ae_checks.sas',
        '06:00',
        'Daily',
        '',
        '',
        'High'
    ])
    ws_sched.append([
        'GZME',
        'Z:\\qa\\ly3537031\\j2s_mc_gzme\\diva\\programs\\oversight\\diva',
        'fmcall.sas',
        '09:00',
        'Weekly',
        'Tue, Thu',
        '2026-12-31',
        'Medium'
    ])

    # Set column widths
    widths = [15, 60, 30, 12, 12, 20, 15, 10]
    for i, w in enumerate(widths, 1):
        ws_sched.column_dimensions[get_column_letter(i)].width = w

    # Save to buffer
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
